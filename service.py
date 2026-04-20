from pathlib import Path

from openpyxl import load_workbook

from app.anonymizer.service import AnonymisationService
from app.processors.base import ProcessedArtifact
from app.utils.exporters import text_to_docx, text_to_pdf


class XlsxProcessor:
    def process(self, input_path: Path, output_dir: Path, service: AnonymisationService) -> ProcessedArtifact:
        workbook = load_workbook(filename=str(input_path))
        preview_lines: list[str] = []

        for sheet in workbook.worksheets:
            sheet.title = service.anonymize_text(sheet.title).anonymized_text[:31]
            for row in sheet.iter_rows():
                rendered = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value and not value.startswith("="):
                        anonymized = service.anonymize_text(value).anonymized_text
                        cell.value = anonymized
                        rendered.append(anonymized)
                    elif value is not None:
                        rendered.append(str(value))
                if rendered:
                    preview_lines.append(" | ".join(rendered))

        stem = input_path.stem + "_anonymised"
        xlsx_path = output_dir / f"{stem}.xlsx"
        docx_path = output_dir / f"{stem}.docx"
        pdf_path = output_dir / f"{stem}.pdf"
        workbook.save(xlsx_path)
        preview_text = "\n".join(preview_lines)
        text_to_docx(preview_text, docx_path, title=input_path.name)
        text_to_pdf(preview_text, pdf_path, title=input_path.name)
        return ProcessedArtifact(input_path=input_path, output_files=[xlsx_path, docx_path, pdf_path], preview_text=preview_text[:3000])
